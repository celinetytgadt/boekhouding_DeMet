# maak-excel.py — zet jaar 3 (oefendossier 1e kwartaal 2028) in de Excel
# van de case Fietsatelier Noor, naast jaar 1 en jaar 2.
#
#   python3 hulpmiddelen/noor/maak-excel.py "case fietsatelier noor.xlsx" "uit.xlsx"
#
# Leest hulpmiddelen/noor/case-j3.json (maak dat eerst met case-j3.js) en
# voegt deze tabbladen toe (bestaan ze al, dan worden ze vervangen):
#   J3 Documenten, J3 Lijnen, J3 Lonen, J3 Uittreksels, J3 Journaal,
#   J3 Proefbalans, J3 Voorraad, J3 Afsluiting
# Daarnaast: de nieuwe bedrijven in Bedrijven, jaar 3 in Overzicht sets en
# een alinea in de Leeswijzer. Jaar 1 en jaar 2 blijven onaangeroerd.
#
# Nodig: Python 3 met openpyxl (pip install openpyxl).

import json
import os
import sys
from copy import copy
from datetime import date

import openpyxl
from openpyxl.utils import get_column_letter

HIER = os.path.dirname(os.path.abspath(__file__))
CASE = json.load(open(os.path.join(HIER, "case-j3.json"), encoding="utf-8"))

if len(sys.argv) != 3:
    sys.exit('Gebruik: python3 maak-excel.py "case fietsatelier noor.xlsx" "uit.xlsx"')
wb = openpyxl.load_workbook(sys.argv[1])

# De opmaak van de kopregel nemen we over van jaar 2.
VOORBEELD = wb["J2 Documenten"]["A1"]
KLEUR_LES = ["FFDDEBF7", "FFE2EFDA", "FFFCE4D6", "FFFFF2CC", "FFE4DFEC", "FFD9D9D9"]


def d(iso):
    return date.fromisoformat(iso) if iso else None


def blad(naam, kop, rijen, breedtes=None, kleurKolom=None):
    if naam in wb.sheetnames:
        del wb[naam]
    ws = wb.create_sheet(naam)
    ws.append(kop)
    for c in ws[1]:
        c.font = copy(VOORBEELD.font)
        c.fill = copy(VOORBEELD.fill)
        c.alignment = copy(VOORBEELD.alignment)
        c.border = copy(VOORBEELD.border)
    for r in rijen:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, b in enumerate(breedtes or []):
        ws.column_dimensions[get_column_letter(i + 1)].width = b
    for rij in ws.iter_rows(min_row=2):
        for c in rij:
            if isinstance(c.value, date):
                c.number_format = "DD/MM/YYYY"
            elif isinstance(c.value, float) or isinstance(c.value, int):
                c.number_format = "#,##0.00;-#,##0.00;0"
    if kleurKolom is not None:
        lessen = []
        for rij in ws.iter_rows(min_row=2):
            les = rij[kleurKolom].value
            if les not in lessen:
                lessen.append(les)
            kleur = KLEUR_LES[lessen.index(les) % len(KLEUR_LES)]
            fill = openpyxl.styles.PatternFill("solid", fgColor=kleur)
            for c in rij:
                c.fill = fill
    return ws


docs = CASE["documenten"]
per_id = {x["id"]: x for x in docs}

# Waar werd elk document betaald? (uittreksel/lijn)
betaald = {}
for u in CASE["uittreksels"]:
    for id_ in filter(None, [s.strip() for s in u["verwijst"].split(",")]):
        betaald.setdefault(id_, []).append("BANK%02d lijn %d (%s)" % (u["uittreksel"], u["lijn"], u["datum"][8:10] + "/" + u["datum"][5:7]))

# ---------------- J3 Documenten ----------------
rijen = []
for x in docs:
    opm = []
    if x.get("fk"):
        opm.append("financiële korting " + ("binnen de termijn betaald" if x.get("fkBetaald") else "te laat betaald, geen korting"))
    if x.get("periode") and x["soort"] in ("AK",):
        opm.append("over te dragen kosten: periode " + x["periode"])
    rijen.append([
        x["ref"], x["id"], x["categorie"], x["soort"], d(x.get("datum")), d(x.get("vervaldag")), str(x.get("nr", "")),
        x.get("van"), x.get("aan"), x.get("verwijst"), x.get("reden") or x.get("periode") or x.get("boekOms"),
        x.get("mvh21"), x.get("btw21"), x.get("mvh6"), x.get("btw6"), x.get("vrij"), x.get("waarborg"), x.get("fk") or None,
        x.get("totaal"), x.get("binnenTermijn"), x.get("contant") if x["soort"] in ("DO", "KAS") else None,
        x.get("bancontact"), "; ".join(betaald.get(x["id"], [])), "; ".join(opm),
    ])
blad("J3 Documenten", ["Ref in de app", "ID", "Les (categorie)", "Soort", "Datum", "Vervaldag", "Documentnr", "Van", "Aan",
                       "Verwijst naar", "Reden / omschrijving", "MvH 21%", "Btw 21%", "MvH 6%", "Btw 6%", "Vrijgesteld",
                       "Waarborg", "Fin. korting", "Totaal", "Te betalen binnen kortingstermijn", "Contant (kas)", "Bancontact",
                       "Betaald via uittreksel", "Opmerking"], rijen,
     [8, 7, 34, 7, 11, 11, 16, 26, 26, 14, 50, 10, 10, 10, 10, 10, 10, 10, 11, 11, 10, 10, 34, 40], kleurKolom=2)

# ---------------- J3 Lijnen ----------------
rijen = []
for x in docs:
    for l in x.get("lijnen", []) if x["soort"] != "BANK" else []:
        rk = l.get("rek")
        rk = ", ".join("%s (%s)" % (r["rek"], r["bedrag"]) for r in rk) if isinstance(rk, list) else rk
        rijen.append([x["id"], x["ref"], str(x.get("nr", "")), l["lijn"], l["soort"], l.get("artikel"), l["oms"], l.get("aantal"),
                      l.get("prijs"), l["bedrag"], l["btwPct"], l.get("btw"), rk])
blad("J3 Lijnen", ["ID", "Ref", "Documentnr", "Lijn", "Soort lijn", "Artikel", "Omschrijving", "Aantal", "Eenheidsprijs",
                   "Bedrag excl. btw", "Btw %", "Btw", "Rekening (vrije lijn)"], rijen,
     [8, 7, 16, 5, 16, 8, 60, 8, 12, 12, 7, 10, 30])

# ---------------- J3 Lonen ----------------
rijen = []
for x in docs:
    if x["soort"] != "LOON":
        continue
    for w in x["werknemers"]:
        rijen.append([x["ref"], x["nr"], x["statuut"], x["periode"], w["naam"], w["loon"], w.get("extra") or None,
                      w.get("extraOms"), w["bruto"], w["rsz"], w["belastbaar"], w["bv"], w["netto"], w["wg"], w["bruto"] + w["wg"]])
blad("J3 Lonen", ["Ref", "Loonstaat", "Statuut", "Periode", "Werknemer", "Basisloon", "Bijkomend", "Wat", "Brutoloon",
                  "RSZ werknemer 13,07%", "Belastbaar", "Bedrijfsvoorheffing", "Netto", "RSZ werkgever 25%", "Totale loonkost"],
     rijen, [8, 18, 10, 14, 38, 10, 10, 16, 10, 12, 11, 12, 10, 12, 12])

# ---------------- J3 Uittreksels ----------------
bb_bank = next(r["saldo"] for r in CASE["beginbalans"] if r["rekening"] == "550000")
rijen = [[None, None, None, None, None, "Saldo bank bij de start van het jaar", None, None, bb_bank]]
for u in CASE["uittreksels"]:
    rijen.append(["BANK%02d" % u["uittreksel"], u["uittreksel"], d(u["datumUittreksel"]), d(u["datum"]), u["lijn"], u["oms"],
                  u["verwijst"], u["bedrag"], u["saldo"]])
blad("J3 Uittreksels", ["Ref", "Uittreksel", "Datum uittreksel", "Datum lijn", "Lijn", "Omschrijving op het uittreksel",
                        "Verwijst naar (ID)", "Bedrag", "Saldo na lijn"], rijen, [8, 9, 12, 12, 5, 80, 18, 11, 11])

# ---------------- J3 Journaal ----------------
rijen = []
for i, j in enumerate(CASE["journaal"]):
    x = per_id[j["doc"]]
    rijen.append([i + 1, j["ref"], j["doc"], d(j["datum"]), x["categorie"], j["rekening"], j["naam"], j["debet"] or None,
                  j["credit"] or None, j.get("relatie"), j["oms"]])
tot_d = round(sum(j["debet"] for j in CASE["journaal"]), 2)
tot_c = round(sum(j["credit"] for j in CASE["journaal"]), 2)
rijen.append([None, None, None, None, None, None, "TOTAAL (zonder beginbalans)", tot_d, tot_c])
blad("J3 Journaal", ["Volgnr", "Ref", "Document", "Datum", "Les", "Rekening", "Rekeningnaam", "Debet", "Credit", "Relatie",
                     "Omschrijving"], rijen, [7, 9, 8, 11, 34, 9, 40, 11, 11, 26, 70], kleurKolom=4)

# ---------------- J3 Proefbalans ----------------
rijen = [[r["rekening"], r["naam"], r["apko"], r["opening"], r["slot"]] for r in CASE["proefbalans"]]
rijen.append([None, "TOTAAL (moet 0 zijn)", None, round(sum(r["opening"] for r in CASE["proefbalans"]), 2),
              round(sum(r["slot"] for r in CASE["proefbalans"]), 2)])
blad("J3 Proefbalans", ["Rekening", "Naam", "A/P/K/O", "Beginbalans 1/1/2028", "Saldo 31/3/2028 na afsluiting",
                        "Debetsaldo = positief, creditsaldo = negatief"], rijen, [10, 50, 8, 18, 22, 40])

# ---------------- J3 Voorraad ----------------
rijen = [[v["artikel"], v["oms"], v["begin"], v["aankopen"], v["retoursAankoop"], v["verkopen"], v["retoursVerkoop"],
          v["dagontvangsten"], v["inventaris"], v["prijs"], v["waarde"]] for v in CASE["voorraad"]]
rijen.append([None, "Totaal", None, None, None, None, None, None, None, None, round(sum(v["waarde"] for v in CASE["voorraad"]), 2)])
blad("J3 Voorraad", ["Artikel", "Omschrijving", "Begin", "+ aankopen", "- retours aankoop", "- verkopen (factuur)",
                     "+ retours verkoop", "- dagontvangsten", "Inventaris 31/3", "Aankoopprijs", "Waarde inventaris"],
     rijen, [8, 34, 7, 10, 10, 12, 10, 12, 11, 11, 14])

# ---------------- J3 Afsluiting ----------------
res = CASE["resultaat"]
rijen = [["AFSCHRIJVINGEN 1E KWARTAAL"], ["Omschrijving", "Rekening", "Aanschaf", "Jaren", "Per jaar", "Reeds afgeschreven",
                                          "1e kwartaal 2028", "Boekwaarde 31/3"]]
for a in CASE["activa"]:
    rijen.append([a["oms"], a["rekening"], a["waarde"], a["jaren"], a["jaar"], a["reeds"], a["kwartaal"], a["boek"]])
rijen += [[], ["LENING " + CASE["lening"]["nr"]], ["Nr", "Vervaldag", "Kapitaal", "Intrest", "Te betalen", "Nog af te lossen"]]
for r in CASE["lening"]["tabel"]:
    rijen.append([r["nr"], d(r["datum"]), r["kapitaal"], r["intrest"], r["totaal"], r["rest"]])
rijen += [[], ["RESULTAAT"],
          ["Opbrengsten", res["opbrengsten"]], ["Kosten (vóór belasting)", res["kosten"]], ["Winst vóór belasting", res["winst"]],
          ["Vennootschapsbelasting 20%", res["belasting"]], ["Winst na belasting", res["restwinst"]],
          ["Wettelijke reserve (5%, tot 10% van het kapitaal)", res["reserve"]], ["Over te dragen winst", res["overdragen"]],
          [], ["OPEN OP 31 MAART 2028"]]
for o in CASE["openstaand"]:
    rijen.append([o["relatie"], o["rekening"], o["saldo"]])
ws = blad("J3 Afsluiting", ["Afsluiting 31 maart 2028 (zie ook J3 Journaal: BTW, AFSCHR, LENING, VR, BELASTING, RESULTAAT)"],
          rijen, [48, 12, 12, 10, 10, 18, 16, 16])
for rij in ws.iter_rows(min_row=2):
    if rij[0].value and isinstance(rij[0].value, str) and rij[0].value.isupper():
        rij[0].font = openpyxl.styles.Font(bold=True)

# ---------------- Bedrijven ----------------
ws = wb["Bedrijven"]
bekend = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
for b in CASE["bedrijven"]:
    if b["code"] not in bekend:
        ws.append([b["code"], b["naam"], b["rol"], b["adres"], b["gemeente"], b["btw"], b["iban"], b["wat"] + " (vanaf jaar 3)"])

# ---------------- Overzicht sets ----------------
ws = wb["Overzicht sets"]
for r in range(ws.max_row, 1, -1):
    if ws.cell(r, 1).value == 3:
        ws.delete_rows(r)
tel = {}
for x in docs:
    t = tel.setdefault(x["categorie"], {"AK": 0, "CN-AK": 0, "VK": 0, "CN-VK": 0, "DO": 0, "KAS": 0, "MEMO": 0, "BANK": 0, "LOON": 0})
    s = {"RSZ": "AK", "AFSCHR": "MEMO", "VR": "MEMO", "LENING": "MEMO"}.get(x["soort"], x["soort"])
    t[s] = t.get(s, 0) + 1
for cat, t in tel.items():
    extra = "loonstaten: %d" % t["LOON"] if t["LOON"] else ""
    ws.append([3, cat, t["AK"], t["CN-AK"], t["VK"], t["CN-VK"], t["DO"], t["KAS"], t["MEMO"], t["BANK"], extra])

# ---------------- Leeswijzer ----------------
ws = wb["Leeswijzer"]
if not any(ws.cell(r, 1).value == "Jaar 3 = 2028 (3e graad, oefenen)" for r in range(1, ws.max_row + 1)):
    ws.insert_rows(5, 3)
    ws.cell(5, 1).value = "Jaar 3 = 2028 (3e graad, oefenen)"
    ws.cell(5, 2).value = (
        "Oefendossier in de boekhoudapp van kern 8 (noor.html): het 1e kwartaal 2028 met een afsluiting op 31 maart, zoals MetWear. "
        "Beginbalans = eindbalans 2027 na afsluiting. Noor heeft nu personeel (Jonas, arbeider; Lina, bediende deeltijds), "
        "een sociaal secretariaat (Loonwijs), een verzekeraar (Gerzo) en een opslagruimte (Immo Vaartkom). "
        "De stukken staan per les (5 à 6 per les): kortingen · kosten, verpakking, investeringen · creditnota's en over te dragen kosten · "
        "verkopen (2 lessen) · loonstaten · RSZ en personeelskosten · financiële verrichtingen · btw, eindejaar en resultaat. "
        "Tabbladen J3 …; gemaakt door hulpmiddelen/noor/ in de repo boekhouding_DeMet.")
    ws.cell(6, 1).value = "Jaar 4 (3e graad, consolidatie)"
    ws.cell(6, 2).value = "Nog te maken: alle stukken door elkaar, als eindopdracht (vanaf volgend schooljaar in de plaats van MetWear)."
    ws.cell(7, 1).value = "Afspraken jaar 3"
    ws.cell(7, 2).value = (
        "MAR van de 3e graad (js/mar.js) plus 700000 voor herstellingen. Over te dragen kosten worden bij de aankoop meteen gesplitst "
        "(kost tot 31/3, rest op 490000). Lonen: RSZ werknemer 13,07%, werkgever 25%; RSZ en voorheffing via factuur van Loonwijs (440000). "
        "Afschrijvingen: een kwart van de jaarafschrijving. Lening: deel dat binnen 12 maanden vervalt naar 423000. "
        "Belasting 20%, wettelijke reserve 5% van de winst na belasting.")
    for r in (5, 6, 7):
        ws.cell(r, 1).font = copy(ws.cell(3, 1).font)
        ws.cell(r, 2).alignment = copy(ws.cell(3, 2).alignment)

wb.save(sys.argv[2])
print("Geschreven: " + sys.argv[2])
